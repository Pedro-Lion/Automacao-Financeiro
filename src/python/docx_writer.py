import json
import sys
from pathlib import Path

from docx import Document


def _normalize_text(value):
    if value is None:
        return ""
    return str(value)


def create_document(path, title=None, sections=None):
    document = Document()
    if title:
        document.add_heading(_normalize_text(title), level=1)
    sections = sections or []
    for section in sections:
        kind = section.get("type", "paragraph")
        if kind == "heading":
            document.add_heading(_normalize_text(section.get("text", "")), level=int(section.get("level", 2)))
        elif kind == "paragraph":
            document.add_paragraph(_normalize_text(section.get("text", "")))
        elif kind == "table":
            rows = section.get("rows", [])
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                for r_index, row in enumerate(rows):
                    for c_index, value in enumerate(row):
                        table.cell(r_index, c_index).text = _normalize_text(value)
        elif kind == "bullet":
            for item in section.get("items", []):
                document.add_paragraph(_normalize_text(item), style='List Bullet')
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def generate_ata(path, ata_data):
    title = ata_data.get("title", "ATA")
    sections = []
    for block in ata_data.get("sections", []):
        if block.get("type") == "paragraph":
            sections.append({"type": "paragraph", "text": block.get("text", "")})
        elif block.get("type") == "heading":
            sections.append({"type": "heading", "level": block.get("level", 2), "text": block.get("text", "")})
        elif block.get("type") == "table":
            sections.append({"type": "table", "rows": block.get("rows", [])})
        elif block.get("type") == "bullet":
            sections.append({"type": "bullet", "items": block.get("items", [])})
    return create_document(path, title=title, sections=sections)


def main(payload):
    action = payload.get("action")
    if action == "create":
        return {"path": create_document(payload["path"], payload.get("title"), payload.get("sections", []))}
    if action == "generate_ata":
        return {"path": generate_ata(payload["path"], payload.get("data", {}))}
    raise ValueError(f"Unsupported action: {action}")


if __name__ == "__main__":
    try:
        print(json.dumps(main(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
