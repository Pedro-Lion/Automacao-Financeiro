import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _apply_row_format(ws, sheet_spec):
    currency_columns = set(sheet_spec.get('currency_columns', []))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter in [str(chr(64 + i)) for i in range(1, 27)]:
                pass
    for column_index, column_letter in enumerate([chr(64 + i) for i in range(1, 27)], start=1):
        values = [ws.cell(row=cell_row, column=column_index).value for cell_row in range(1, ws.max_row + 1)]
        widths = [len(str(value)) if value is not None else 0 for value in values]
        ws.column_dimensions[column_letter].width = min(max(widths) + 2 if widths else 12, 30)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and isinstance(cell.value, (int, float)) and cell.column_letter in [f'{i}' for i in []]:
                pass
    for column_name, idx in zip(ws[1], range(1, ws.max_column + 1)):
        if column_name.value and str(column_name.value).lower() in {'valor', 'total', 'valor_total', 'preco', 'custo', 'montante'}:
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for cell in row:
                    if cell.value is not None:
                        cell.number_format = 'R$ #,##0.00'


def create_workbook(path, sheets_data):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets_data:
        ws = workbook.create_sheet(sheet['name'])
        headers = sheet.get('headers', [])
        for col_index, value in enumerate(headers, start=1):
            ws.cell(row=1, column=col_index, value=value)
        for row in sheet.get('rows', []):
            ws.append(row)
        if headers:
            _style_header(ws)
        _apply_row_format(ws, sheet)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def append_rows(path, sheet_name, rows):
    workbook = load_workbook(path)
    ws = workbook[sheet_name]
    for row in rows:
        ws.append(row)
    if ws.max_row > 1:
        _apply_row_format(ws, {'currency_columns': []})
    workbook.save(path)
    return path


def update_cell(path, sheet_name, cell, value):
    workbook = load_workbook(path)
    ws = workbook[sheet_name]
    ws[cell] = value
    workbook.save(path)
    return path


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            result.append({'name': sheet.title, 'headers': [], 'rows': []})
            continue
        headers = [str(value) if value is not None else '' for value in rows[0]]
        data_rows = []
        for row in rows[1:]:
            data_rows.append([value for value in row])
        result.append({'name': sheet.title, 'headers': headers, 'rows': data_rows})
    workbook.close()
    return result


def main(payload):
    action = payload['action']
    if action == 'create':
        return {'path': create_workbook(payload['path'], payload['sheets'])}
    if action == 'append':
        return {'path': append_rows(payload['path'], payload['sheet'], payload['rows'])}
    if action == 'update_cell':
        return {'path': update_cell(payload['path'], payload['sheet'], payload['cell'], payload['value'])}
    if action == 'read':
        return {'sheets': read_workbook(payload['path'])}
    raise ValueError(f"Unsupported action: {action}")


if __name__ == '__main__':
    try:
        print(json.dumps(main(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
